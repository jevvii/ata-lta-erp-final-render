import os
import sys
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ['sz', 'val', 'color', 'space']:
                if key in edge_data:
                    element.set(qn(f'w:{key}'), str(edge_data[key]))

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement('w:cantSplit')
    trPr.append(cantSplit)

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    trPr.append(tblHeader)

def format_inline_runs(paragraph, text, base_font="Calibri", base_size=10.5, base_color=(31, 41, 55), default_bold=False):
    pattern = re.compile(r'(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))')
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            inner = token[2:-2]
            run = paragraph.add_run(inner)
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.bold = True
            run.font.color.rgb = RGBColor(*base_color)
        elif token.startswith('*') and token.endswith('*'):
            inner = token[1:-1]
            run = paragraph.add_run(inner)
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.italic = True
            run.font.color.rgb = RGBColor(*base_color)
        elif token.startswith('`') and token.endswith('`'):
            inner = token[1:-1]
            run = paragraph.add_run(inner)
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 0.5)
            run.font.color.rgb = RGBColor(15, 23, 42)
            run.font.bold = default_bold
        elif token.startswith('[') and ']' in token and '(' in token and token.endswith(')'):
            m = re.match(r'\[([^\]]+?)\]\(([^)]+?)\)', token)
            if m:
                label, url = m.groups()
                run = paragraph.add_run(label)
                run.font.name = base_font
                run.font.size = Pt(base_size)
                run.font.color.rgb = RGBColor(37, 99, 235)
                run.underline = True
            else:
                run = paragraph.add_run(token)
                run.font.name = base_font
                run.font.size = Pt(base_size)
                run.font.color.rgb = RGBColor(*base_color)
        else:
            run = paragraph.add_run(token)
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.bold = default_bold
            run.font.color.rgb = RGBColor(*base_color)

def convert_laymans_markdown_to_docx(md_path, docx_path, img_dir):
    print(f"Reading {md_path}...")
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()

    # Configure margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("ATA & LTA ERP Platform — Visual User Guide & Manual (Layman's Edition)")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun1 = fp.add_run("CONFIDENTIAL — INTERNAL OPERATIONAL USE ONLY\t\tVisual Reference Manual")
        frun1.font.name = "Calibri"
        frun1.font.size = Pt(8.5)
        frun1.font.color.rgb = RGBColor(148, 163, 184)

    # Base style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(31, 41, 55)

    i = 0
    total_lines = len(lines)
    in_code_block = False
    code_block_lines = []
    code_block_lang = ""

    while i < total_lines:
        line = lines[i].rstrip('\r\n')

        # Code block
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_block_lang = line[3:].strip()
                code_block_lines = []
                i += 1
                continue
            else:
                in_code_block = False
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                cell.width = Inches(6.5)
                set_cell_shading(cell, "F8FAFC")
                set_cell_border(cell, 
                                left={'sz': 16, 'val': 'single', 'color': '2563EB'},
                                top={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                bottom={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                right={'sz': 4, 'val': 'single', 'color': 'E2E8F0'})
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(2)
                r_body = cp.add_run("\n".join(code_block_lines))
                r_body.font.name = "Consolas"
                r_body.font.size = Pt(9.0)
                r_body.font.color.rgb = RGBColor(30, 41, 59)
                i += 1
                continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Markdown Images: ![Caption](path)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', line.strip())
        if img_match:
            caption, img_rel_path = img_match.groups()
            filename = os.path.basename(img_rel_path)
            full_img_path = os.path.join(img_dir, filename)
            
            if os.path.exists(full_img_path):
                # Add picture
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.paragraph_format.keep_with_next = True
                run_img = p_img.add_run()
                run_img.add_picture(full_img_path, width=Inches(6.0))
                
                # Add caption below picture
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(2)
                p_cap.paragraph_format.space_after = Pt(8)
                run_cap = p_cap.add_run(f"📸 {caption}")
                run_cap.font.name = "Calibri"
                run_cap.font.size = Pt(9.0)
                run_cap.font.italic = True
                run_cap.font.bold = True
                run_cap.font.color.rgb = RGBColor(71, 85, 105)
            else:
                p_missing = doc.add_paragraph(f"[Image: {caption} - {filename}]")
                p_missing.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # Markdown Tables: | col1 | col2 | ...
        if line.strip().startswith('|') and line.strip().endswith('|'):
            table_lines = []
            while i < total_lines and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            if len(table_lines) >= 2:
                raw_rows = []
                for tline in table_lines:
                    parts = [p.strip() for p in tline.split('|')[1:-1]]
                    raw_rows.append(parts)
                
                header_row = raw_rows[0]
                data_rows = []
                if len(raw_rows) > 1 and all(set(c).issubset({'-', ':', ' '}) for c in raw_rows[1]):
                    data_rows = raw_rows[2:]
                else:
                    data_rows = raw_rows[1:]
                
                num_cols = len(header_row)
                if num_cols > 0:
                    tbl = doc.add_table(rows=len(data_rows) + 1, cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    
                    # Style Header
                    h_row = tbl.rows[0]
                    set_repeat_header(h_row)
                    prevent_row_split(h_row)
                    for c_idx, cell in enumerate(h_row.cells):
                        if c_idx < len(header_row):
                            cell.text = ""
                            cp = cell.paragraphs[0]
                            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            cp.paragraph_format.space_before = Pt(4)
                            cp.paragraph_format.space_after = Pt(4)
                            format_inline_runs(cp, header_row[c_idx], base_font="Calibri", base_size=9.5, base_color=(255, 255, 255), default_bold=True)
                        set_cell_shading(cell, "1E3A8A")
                        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                        set_cell_border(cell,
                                        top={'sz': 6, 'val': 'single', 'color': '1E3A8A'},
                                        bottom={'sz': 6, 'val': 'single', 'color': '1E3A8A'},
                                        left={'sz': 4, 'val': 'single', 'color': 'CBD5E1'},
                                        right={'sz': 4, 'val': 'single', 'color': 'CBD5E1'})
                    
                    # Style Data Rows
                    for r_idx, d_row in enumerate(data_rows):
                        t_row = tbl.rows[r_idx + 1]
                        prevent_row_split(t_row)
                        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
                        for c_idx, cell in enumerate(t_row.cells):
                            if c_idx < len(d_row):
                                cell.text = ""
                                cp = cell.paragraphs[0]
                                cp.paragraph_format.space_before = Pt(3)
                                cp.paragraph_format.space_after = Pt(3)
                                cell_val = d_row[c_idx]
                                if cell_val.strip() in ['✅', '❌', 'Yes', 'No']:
                                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                else:
                                    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                                format_inline_runs(cp, cell_val, base_font="Calibri", base_size=9.0, base_color=(31, 41, 55))
                            
                            set_cell_shading(cell, bg_color)
                            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                            set_cell_border(cell,
                                            top={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                            bottom={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                            left={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                            right={'sz': 4, 'val': 'single', 'color': 'E2E8F0'})
                    
                    sp = doc.add_paragraph()
                    sp.paragraph_format.space_before = Pt(0)
                    sp.paragraph_format.space_after = Pt(6)
            continue

        # Horizontal Divider (---)
        if line.strip() == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run("―" * 55)
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(203, 213, 225)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Title (# )
        if line.startswith('# '):
            text = line[2:].strip()
            if re.match(r'^[0-9]+\.\s', text):
                doc.add_page_break()
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(20)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(text)
                run.font.name = "Calibri"
                run.font.size = Pt(18)
                run.font.bold = True
                run.font.color.rgb = RGBColor(30, 58, 138)
            elif text == "Table of Contents":
                doc.add_page_break()
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(text)
                run.font.name = "Calibri"
                run.font.size = Pt(16)
                run.font.bold = True
                run.font.color.rgb = RGBColor(30, 58, 138)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(24)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(text)
                run.font.name = "Calibri"
                run.font.size = Pt(22)
                run.font.bold = True
                run.font.color.rgb = RGBColor(30, 58, 138)
            i += 1
            continue

        # Subtitle (## )
        if line.startswith('## '):
            text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        # Section Level 3 (### )
        if line.startswith('### '):
            text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(37, 99, 235)
            i += 1
            continue

        # Section Level 4 (#### )
        if line.startswith('#### '):
            text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(55, 65, 81)
            i += 1
            continue

        # Bullet items: * or -
        if line.strip().startswith('* ') or line.strip().startswith('- '):
            raw_text = line.strip()[2:].strip()
            indent_spaces = len(line) - len(line.lstrip())
            level = 1 if indent_spaces >= 2 else 0
            
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            if level == 1:
                p.paragraph_format.left_indent = Inches(0.5)
            
            format_inline_runs(p, raw_text, base_font="Calibri", base_size=10.0, base_color=(31, 41, 55))
            i += 1
            continue

        # Numbered list items: 1. , 2. 
        m_num = re.match(r'^(\s*)([0-9]+)\.\s+(.*)$', line)
        if m_num:
            indent_spaces = len(m_num.group(1))
            num_str = m_num.group(2)
            content_str = m_num.group(3)
            level = 1 if indent_spaces >= 2 else 0
            
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Inches(0.5 if level == 1 else 0.25)
            
            r_num = p.add_run(f"{num_str}.  ")
            r_num.font.name = "Calibri"
            r_num.font.size = Pt(10.0)
            r_num.font.bold = True
            r_num.font.color.rgb = RGBColor(30, 58, 138)
            
            format_inline_runs(p, content_str, base_font="Calibri", base_size=10.0, base_color=(31, 41, 55))
            i += 1
            continue

        # Blockquote (> )
        if line.strip().startswith('> '):
            quote_text = line.strip()[2:].strip()
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            cell.width = Inches(6.5)
            set_cell_shading(cell, "F1F5F9")
            set_cell_border(cell,
                            left={'sz': 20, 'val': 'single', 'color': '2563EB'},
                            top={'sz': 0, 'val': 'none', 'color': 'auto'},
                            bottom={'sz': 0, 'val': 'none', 'color': 'auto'},
                            right={'sz': 0, 'val': 'none', 'color': 'auto'})
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(2)
            format_inline_runs(cp, quote_text, base_font="Calibri", base_size=10.0, base_color=(51, 65, 85))
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        format_inline_runs(p, line, base_font="Calibri", base_size=10.5, base_color=(31, 41, 55))
        i += 1

    print(f"Saving formatted Word document with screenshots to {docx_path}...")
    doc.save(docx_path)
    print("DOCX successfully generated!")

print("DOCX conversion engine ready.")
